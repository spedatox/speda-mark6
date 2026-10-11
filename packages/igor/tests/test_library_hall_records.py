"""Behavior regressions for real Hall indexing/retrieval, isolated from owner data."""
import copy
import hashlib
import io
import uuid
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, HTTPException
from httpx import ASGITransport, AsyncClient
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.context import AgentContext
from app.database import Base
from app.models.background_job import BackgroundJob
from app.models.library import (
    LibrarySource, LibraryDocument, LibraryRevision, LibraryIndexGeneration,
    LibraryPassage, LibraryEmbedding, LibraryCitation,
)
from app.services.library import LibraryService
from app.services.library_extract import ExtractionError, chunk_text, extract_document_bytes
from app.services.library_indexer import LibraryIndexer
from app.skills.library import ListLibrarySkill, ReadLibrarySkill, SearchLibrarySkill


class Catalog:
    hisar_url = "http://isolated-hisar"
    max_bytes = 25 * 1024**2
    configured = True

    def __init__(self):
        self.records, self.blobs, self.events, self.receipts, self.queued = {}, {}, [], [], {}
        self.fail_queue = False
        self.bad_bytes = False

    def add(self, title, content, *, doc_id=None, mime="text/plain", **edition):
        doc_id = doc_id or str(uuid.uuid4())
        revision_id = str(uuid.uuid4())
        blob = content.encode() if isinstance(content, str) else content
        if doc_id not in self.records:
            self.records[doc_id] = {"document": {
                "id": doc_id, "title": title, "aliases": [], "category": "Academics", "type": "reference",
                "institution": None, "subject": None, "active_revision_id": None, "access_policy": {},
                "is_archived": False, "is_tombstone": False, "metadata_version": 1,
                "created_at": "2026-10-10T10:00:00+00:00", "updated_at": "2026-10-10T10:00:00+00:00",
            }, "revisions": []}
        record = self.records[doc_id]
        rev = {"id": revision_id, "document_id": doc_id, "sequence": len(record["revisions"])+1,
               "blob_sha256": hashlib.sha256(blob).hexdigest(), "original_filename": "original.txt",
               "mime_type": mime, "size_bytes": len(blob), "issue_date": None, "coverage_start": None,
               "coverage_end": None, "term": None, "superseded_by": None, "created_at": "2026-10-10T10:00:00+00:00",
               **edition}
        record["revisions"].append(rev)
        record["active_revision"] = rev
        record["document"]["active_revision_id"] = revision_id
        record["document"]["metadata_version"] += 1
        self.blobs[revision_id] = blob
        self.events.append({"sequence": len(self.events)+1, "event_type": "revision_added",
                            "document_id": doc_id, "revision_id": revision_id})
        return doc_id, revision_id

    async def get_document(self, doc_id):
        return copy.deepcopy(self.records[doc_id])

    async def get_manifest(self, offset=0, limit=100):
        records = list(self.records.values())
        return {"records": copy.deepcopy(records[offset:offset+limit]), "has_more": offset+limit < len(records),
                "next_offset": offset+limit}

    async def get_events(self, after_sequence=0):
        events = self.events[after_sequence:after_sequence+2]
        return {"events": copy.deepcopy(events), "has_more": after_sequence+len(events) < len(self.events)}

    async def resolve_access(self, document_ids, agent_id):
        items = []
        for doc_id in document_ids:
            record = self.records.get(doc_id)
            doc = record["document"] if record else None
            agents = doc["access_policy"].get("allowed_agents") if doc else []
            granted = bool(doc and not doc["is_tombstone"] and (agents is None or agent_id in agents))
            items.append({"document_id": doc_id, "active_revision_id": doc["active_revision_id"] if granted else None,
                          "is_archived": doc["is_archived"] if granted else False, "access_granted": granted,
                          "record": copy.deepcopy(record) if granted else None})
        return {"items": items}

    async def download_revision_blob(self, revision_id):
        return b"tampered" if self.bad_bytes else self.blobs[revision_id]

    async def report_index_receipt(self, document_id, revision_id, status):
        self.receipts.append((document_id, revision_id, status))

    async def enqueue(self, **job):
        if self.fail_queue:
            raise ConnectionError("queue unavailable")
        key = job["unique_key"]
        if key in self.queued:
            return None
        self.queued[key] = job
        return len(self.queued)


@pytest.fixture
async def hall(tmp_path):
    engine = create_async_engine("sqlite+aiosqlite:///" + (tmp_path / "derived.db").as_posix())
    tables = [model.__table__ for model in (
        LibrarySource, LibraryDocument, LibraryRevision, LibraryIndexGeneration,
        LibraryPassage, LibraryEmbedding, LibraryCitation,
        BackgroundJob,
    )]
    async with engine.begin() as db:
        await db.run_sync(lambda conn: Base.metadata.create_all(conn, tables=tables))
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    client = Catalog()
    indexer = LibraryIndexer(client, sessions, enqueue=client.enqueue)
    context = AgentContext(agent_id="ultron", user_id=1, session_id=1, request_id="hall-test",
                           triggered_by="user", trigger_payload={}, output_mode="respond", model="test",
                           system_prompt="", conversation_history=[], db=None, timezone="Europe/Istanbul")
    service = LibraryService(indexer, n8n_secret="isolated-n8n")
    yield SimpleNamespace(client=client, indexer=indexer, service=service, sessions=sessions, context=context)
    await engine.dispose()


async def index_all(hall):
    await hall.indexer.reconcile(full=True)
    for record in hall.client.records.values():
        for rev in record["revisions"]:
            await hall.indexer.index_revision(record["document"]["id"], rev["id"])


def control_settings(**updates):
    values = dict(n8n_secret="isolated-n8n", library_semantic_enabled=False, library_ocr_enabled=False,
                  library_ocr_languages="tur+eng", library_ocr_timeout_seconds=300,
                  library_embedding_api_key="", openai_api_key="", embedding_model="isolated-model",
                  library_max_pages=200, library_max_bytes=25 * 1024**2)
    return SimpleNamespace(**{**values, **updates})


async def test_settings_are_persisted_live_and_secrets_never_returned(hall, monkeypatch):
    from app.services.library_controls import LibraryControls
    from app.schemas.library import LibrarySettingsUpdate
    saved = []
    settings = control_settings()
    controls = LibraryControls(hall.indexer, settings, writer=saved.append)
    # No network call: test only the authenticated settings transaction here.
    result = await controls.save("isolated-n8n", LibrarySettingsUpdate(
        semantic_enabled=False, ocr_enabled=False, embedding_api_key="test-secret-only"))
    assert result["saved"] and result["embedding_key_configured"]
    assert "test-secret-only" not in str(result)
    assert saved[0]["LIBRARY_EMBEDDING_API_KEY"] == "test-secret-only"
    assert settings.library_embedding_api_key == "test-secret-only"
    assert hall.indexer.embedder is None
    with pytest.raises(HTTPException) as denied:
        await controls.snapshot("wrong")
    assert denied.value.status_code == 401
    with pytest.raises(HTTPException):
        await LibraryControls(hall.indexer, control_settings(), writer=saved.append).save("isolated-n8n",
            LibrarySettingsUpdate(semantic_enabled=True, ocr_enabled=False))
    assert len(saved) == 1
    monkeypatch.setattr(controls, "ocr_runtime", lambda: {"available": False, "languages": []})
    with pytest.raises(HTTPException) as missing:
        await controls.save("isolated-n8n", LibrarySettingsUpdate(semantic_enabled=False, ocr_enabled=True))
    assert missing.value.status_code == 422 and len(saved) == 1


def test_pdf_progress_counts_actual_pages():
    writer = PdfWriter()
    writer.add_blank_page(100, 100); writer.add_blank_page(100, 100)
    blob = io.BytesIO(); writer.write(blob)
    events = []
    class OCR:
        def pdf_page(self, blob, number):
            return f"Read page {number}"
    result = extract_document_bytes(blob.getvalue(), "application/pdf", ocr=OCR(), progress=lambda *event: events.append(event))
    assert len(result) == 2
    assert ("ocr", 0, 2, "pages") in events
    assert events[-1] == ("extracting", 2, 2, "pages")


async def test_progress_tracks_semantic_batches_and_filters_other_owner(hall):
    from app.services.library_controls import LibraryControls
    doc, rev = hall.client.add("Large reference", "Long reference passage. " * 3500)
    await index_all(hall)
    controls = LibraryControls(hall.indexer, control_settings())
    checkpoints = []
    class Embedder:
        fingerprint = "test-batch-progress"
        async def embed(self, texts):
            info = (await controls.progress("isolated-n8n", [doc]))["records"][0]
            checkpoints.append(info["semantic"]["indexed"])
            return [[1.0, 0.0] for _ in texts]
    hall.indexer.embedder = Embedder()
    await hall.indexer.embed_revision(doc, rev)
    assert checkpoints[0] == 0 and checkpoints[1] == 32
    result = await controls.progress("isolated-n8n", [doc, "other-owner-document"])
    info = result["records"][0]
    assert info["progress"]["stage"] == "ready"
    assert info["semantic"]["indexed"] == info["semantic"]["total"] > 32
    assert result["records"][1] == {"document_id": "other-owner-document", "state": "not_synced", "progress": {"stage": "awaiting_sync"}}


async def test_library_drain_does_not_claim_unrelated_jobs(hall, monkeypatch):
    from app.services import task_queue
    import datetime as dt
    async with hall.sessions() as db:
        for kind, owner, source in [("title", 1, ""), ("library_index", 2, hall.indexer.source_id),
                                     ("library_index", 1, "other-source"), ("library_index", 1, hall.indexer.source_id)]:
            db.add(BackgroundJob(kind=kind, user_id=owner, unique_key=str(uuid.uuid4()),
                                 payload={"source_id": source}, status="pending", run_after=dt.datetime.now(dt.timezone.utc)))
        await db.commit()
    monkeypatch.setattr(task_queue, "AsyncSessionLocal", hall.sessions)
    async def run(job_id, *args):
        await task_queue._finish(job_id, error=None)
        return True
    monkeypatch.setattr(task_queue, "_run_one", run)
    await hall.service.drain_jobs()
    async with hall.sessions() as db:
        jobs = list((await db.scalars(select(BackgroundJob).order_by(BackgroundJob.id))).all())
    assert [job.status for job in jobs] == ["pending", "pending", "pending", "done"]


async def test_old_failure_cannot_hide_completed_retry(hall):
    from app.services.library_controls import LibraryControls
    doc, rev = hall.client.add("Reference", "Indexed text")
    await index_all(hall)
    async with hall.sessions() as db:
        for state in ["failed", "done"]:
            db.add(BackgroundJob(kind="library_index", user_id=1, unique_key=hall.indexer.generation_id(rev),
                                 payload={"source_id": hall.indexer.source_id, "revision_id": rev}, status=state))
            await db.commit()
    info = (await LibraryControls(hall.indexer, control_settings()).progress("isolated-n8n", [doc]))["records"][0]
    assert info["state"] == "ready_keyword"


async def test_retry_does_not_reset_running_semantic_work(hall):
    doc, rev = hall.client.add("Reference", "Text ready while semantic work runs")
    await index_all(hall)
    async with hall.sessions() as db:
        db.add(BackgroundJob(kind="library_embed", user_id=1, unique_key="isolated-vector-job",
                             payload={"source_id": hall.indexer.source_id, "revision_id": rev}, status="running"))
        await db.commit()
    assert await hall.indexer._accept(hall.client.records[doc], retry=True) == 0
    async with hall.sessions() as db:
        assert (await db.get(LibraryIndexGeneration, hall.indexer.generation_id(rev))).status == "ready_keyword"


async def test_upload_index_search_read_and_resolve_real_evidence(hall):
    doc, rev = hall.client.add("Akademik takvim", "Final sınavları 18 Ocak tarihinde başlar.")
    await index_all(hall)
    result = await hall.service.search(hall.context, "sinav ne zaman")
    assert result.status == "success"
    assert result.search_mode == "keyword_bm25"
    assert result.results[0].content == "Final sınavları 18 Ocak tarihinde başlar."
    assert result.results[0].revision_id == rev
    read = await hall.service.read_passage(hall.context, doc)
    assert read.citation_ids == [result.results[0].citation_id]
    source = await hall.service.citation_source(read.citation_ids[0])
    assert source["document_id"] == doc and source["revision_id"] == rev
    assert not source["historical"]
    assert (await hall.service.list_documents(hall.context)).documents[0].index_status == "ready_keyword"
    async with hall.sessions() as db:
        assert len((await db.scalars(select(LibraryCitation))).all()) == 1


async def test_pending_replacement_never_falls_back_to_old_edition(hall):
    doc, old = hall.client.add("Calendar", "Final exam starts January 18.")
    await index_all(hall)
    _, new = hall.client.add("Calendar", "Final exam starts January 25.", doc_id=doc)
    # Change authority without delivering its event.
    result = await hall.service.search(hall.context, "exam")
    assert not result.results and result.status == "unavailable"
    assert (await hall.service.read_passage(hall.context, doc)).status == "pending"
    historical = await hall.service.read_passage(hall.context, doc, revision_id=old)
    assert historical.results[0].historical
    await index_all(hall)
    result = await hall.service.search(hall.context, "exam")
    assert all(r.revision_id == new for r in result.results)


async def test_access_filter_precedes_top_k_and_revocation_is_live(hall):
    for n in range(10):
        doc, _ = hall.client.add(f"Denied {n}", "Rareword rareword rareword")
        hall.client.records[doc]["document"]["access_policy"] = {"allowed_agents": ["sentinel"]}
    allowed, _ = hall.client.add("Allowed", "Rareword appears once.")
    await index_all(hall)
    result = await hall.service.search(hall.context, "rareword", limit=1)
    assert result.results[0].document_id == allowed
    hall.client.records[allowed]["document"]["access_policy"] = {"allowed_agents": []}
    assert not (await hall.service.search(hall.context, "rareword")).results
    assert not (await hall.service.read_passage(hall.context, allowed)).results
    assert not (await hall.service.list_documents(hall.context)).documents


async def test_coverage_dates_and_explicit_history(hall):
    old, old_rev = hall.client.add("September menu", "Dinner: rice", coverage_start="2026-09-01", coverage_end="2026-09-30")
    current, _ = hall.client.add("October menu", "Dinner: soup", coverage_start="2026-10-01", coverage_end="2026-10-31")
    await index_all(hall)
    result = await hall.service.search(hall.context, "Dinner", on_date="2026-10-10")
    assert [r.document_id for r in result.results] == [current]
    assert not (await hall.service.search(hall.context, "Dinner", on_date="2026-11-10")).results


async def test_durable_acceptance_cursor_and_replay(hall):
    for n in range(5):
        hall.client.add(str(n), "some text")
    hall.client.fail_queue = True
    with pytest.raises(ConnectionError):
        await hall.indexer.reconcile()
    async with hall.sessions() as db:
        source = await db.get(LibrarySource, hall.indexer.source_id)
        assert source.sync_cursor == 0
    hall.client.fail_queue = False
    first = await hall.indexer.reconcile()
    assert first["cursor"] == 5 and len(hall.client.queued) == 5
    await hall.indexer.reconcile(full=True)
    assert len(hall.client.queued) == 5
    await index_all(hall)
    async with hall.sessions() as db:
        assert len((await db.scalars(select(LibraryIndexGeneration))).all()) == 5


async def test_hash_mismatch_never_publishes_passages(hall):
    doc, rev = hall.client.add("Original", "Important evidence")
    await hall.indexer.reconcile()
    hall.client.bad_bytes = True
    with pytest.raises(ValueError, match="SHA-256"):
        await hall.indexer.index_revision(doc, rev)
    async with hall.sessions() as db:
        assert not (await db.scalars(select(LibraryPassage))).all()
        assert (await db.get(LibraryIndexGeneration, hall.indexer.generation_id(rev))).status == "failed"
    hall.client.bad_bytes = False
    await hall.indexer.index_revision(doc, rev)
    assert (await hall.service.read_passage(hall.context, doc)).status == "success"


async def test_scan_and_unsupported_format_require_review_without_fake_text(hall):
    import io
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    original = io.BytesIO()
    writer.write(original)
    doc, rev = hall.client.add("Scanned", original.getvalue(), mime="application/pdf")
    await index_all(hall)
    assert (await hall.service.read_passage(hall.context, doc)).status == "needs_review"
    async with hall.sessions() as db:
        assert not (await db.scalars(select(LibraryPassage))).all()
    assert hall.client.receipts[-1][-1] == "needs_review"
    with pytest.raises(ExtractionError):
        extract_document_bytes(b"original", "image/png")


async def test_citations_fail_closed_on_withdrawal_or_authority_outage(hall):
    doc, rev = hall.client.add("Original", "Exact evidence")
    await index_all(hall)
    read = await hall.service.read_passage(hall.context, doc)
    hall.client.records[doc]["document"]["is_tombstone"] = True
    with pytest.raises(HTTPException) as withdrawn:
        await hall.service.citation_source(read.citation_ids[0])
    assert withdrawn.value.status_code == 410


async def test_tools_use_registry_contract_and_bound_arguments(hall):
    for tool in (ListLibrarySkill(hall.service), SearchLibrarySkill(hall.service), ReadLibrarySkill(hall.service)):
        assert tool.read_only and not tool.memoize
        assert tool.to_tool_definition()["input_schema"]["type"] == "object"
    assert not SearchLibrarySkill.deferred
    result = await SearchLibrarySkill(hall.service).execute({"query": "x", "limit": 1000}, hall.context)
    assert "invalid_arguments" in result
    other = copy.copy(hall.context)
    other.user_id = 99
    with pytest.raises(HTTPException) as exc:
        await hall.service.list_documents(other)
    assert exc.value.status_code == 403


async def test_reconciliation_probe_requires_both_headers(hall, monkeypatch):
    from app.config import settings
    from app.middleware.auth import AuthMiddleware
    from app.routers.library import router
    monkeypatch.setattr(settings, "speda_api_key", "isolated-api")
    app = FastAPI()
    app.state.library = hall.service
    app.include_router(router)
    app.add_middleware(AuthMiddleware)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://igor") as client:
        assert (await client.post("/library/reconcile")).status_code == 401
        assert (await client.post("/library/reconcile", headers={"X-API-Key": "isolated-api"})).status_code == 401
        # Avoid draining the real owner's database in this isolated test.
        async def no_work():
            pass
        monkeypatch.setattr(hall.service, "drain_jobs", no_work)
        response = await client.post("/library/reconcile", headers={"X-API-Key": "isolated-api", "X-N8N-Secret": "isolated-n8n"})
        assert response.status_code == 200


def test_long_paragraphs_are_bounded_and_binary_errors_are_not_evidence():
    chunks = chunk_text("a" * 10001)
    assert max(map(len, chunks)) <= 2000
    assert "".join(chunks) == "a" * 10001
    with pytest.raises(ExtractionError):
        extract_document_bytes(b"not pdf", "application/pdf")
    with pytest.raises(ExtractionError):
        extract_document_bytes(b"\xff", "text/plain")


def test_text_pdf_preserves_page_locator_and_mixed_scan_requires_review():
    import io
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 30 350 Td (Final exam: January 18) Tj ET")
    page[NameObject("/Contents")] = stream
    original = io.BytesIO()
    writer.write(original)
    passages = extract_document_bytes(original.getvalue(), "application/pdf")
    assert passages[0].page == 1 and "January 18" in passages[0].content
    writer.add_blank_page(width=400, height=400)
    mixed = io.BytesIO()
    writer.write(mixed)
    with pytest.raises(ExtractionError, match="Page 2"):
        extract_document_bytes(mixed.getvalue(), "application/pdf")


async def test_rename_reindexes_title_without_changing_citation_identity(hall):
    doc, rev = hall.client.add("Old title", "Exact source text")
    await index_all(hall)
    before = await hall.service.read_passage(hall.context, doc)
    hall.client.records[doc]["document"]["title"] = "Renamed archive"
    hall.client.records[doc]["document"]["metadata_version"] += 1
    await hall.indexer.reconcile(full=True)
    found = await hall.service.search(hall.context, "Renamed")
    assert found.results[0].document_title == "Renamed archive"
    assert found.results[0].citation_id == before.citation_ids[0]


async def test_legacy_default_source_adoption_keeps_existing_evidence(hall):
    doc, rev = hall.client.add("Retained", "Original evidence")
    await index_all(hall)
    async with hall.sessions() as db:
        db.add(LibrarySource(id="default", hisar_url="http://legacy-alias"))
        await db.flush()
        record = await db.get(LibraryDocument, doc)
        record.source_id = "default"
        await db.commit()
    await hall.indexer.reconcile(full=True)
    async with hall.sessions() as db:
        assert (await db.get(LibraryDocument, doc)).source_id == hall.indexer.source_id
        assert await db.get(LibrarySource, "default")
        assert len((await db.scalars(select(LibraryPassage))).all()) == 1
    assert (await hall.service.read_passage(hall.context, doc)).results[0].content == "Original evidence"


def test_office_extraction_retains_order_tables_and_formula_locators():
    from docx import Document
    from openpyxl import Workbook
    document = Document()
    document.add_heading("Exam dates", level=1)
    document.add_paragraph("Final sınavı 18 Ocak.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Course", "Date"
    table.cell(1, 0).text, table.cell(1, 1).text = "Algorithms", "January 18"
    document.add_paragraph("End of schedule")
    document.sections[0].header.paragraphs[0].text = "Official calendar"
    original = io.BytesIO()
    document.save(original)
    passages = extract_document_bytes(original.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert [p.content for p in passages[:5]] == ["Exam dates", "Final sınavı 18 Ocak.", "Course | Date", "Algorithms | January 18", "End of schedule"]
    assert "Table 1, row 2" in passages[3].section
    assert any(p.content == "Official calendar" and "header" in p.section for p in passages)
    assert all(p.page is None and not p.needs_review for p in passages)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Amount"])
    sheet.append(["Books", 120])
    sheet.append(["Total", "=SUM(B2:B2)"])
    original = io.BytesIO()
    workbook.save(original)
    passages = extract_document_bytes(original.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "A1: Item | B1: Amount" in passages[1].content
    assert "B2: 120" in passages[1].content and "row 2" in passages[1].section
    assert "B3 formula: =SUM(B2:B2)" in passages[2].content
    with pytest.raises(ExtractionError, match="character limit"):
        extract_document_bytes(original.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", max_characters=10)


def test_office_zip_and_visual_evidence_fail_closed():
    import zipfile
    doc_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    original = io.BytesIO()
    with zipfile.ZipFile(original, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b'<!DOCTYPE document [<!ENTITY a "expanded">]><document/>')
    with pytest.raises(ExtractionError, match="entity declarations"):
        extract_document_bytes(original.getvalue(), doc_mime)
    original = io.BytesIO()
    with zipfile.ZipFile(original, "w") as archive:
        archive.writestr("word/media/image.png", b"scan")
    with pytest.raises(ExtractionError, match="visual or embedded"):
        extract_document_bytes(original.getvalue(), doc_mime)


async def test_actual_office_original_indexes_and_resolves_cited_cells(hall):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(["Course", "Final date"])
    workbook.active.append(["Algorithms", "January 18"])
    original = io.BytesIO()
    workbook.save(original)
    doc, rev = hall.client.add("Final timetable", original.getvalue(),
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    await index_all(hall)
    found = await hall.service.search(hall.context, "Algorithms")
    assert found.results[0].revision_id == rev
    assert "row 2" in found.results[0].section and found.results[0].page is None
    source = await hall.service.citation_source(found.results[0].citation_id)
    assert source["content"] == found.results[0].content


class DeterministicEmbeddings:
    fingerprint = "isolated:semantic-v1"

    def __init__(self):
        self.calls = []
        self.fail = False

    async def embed(self, texts):
        self.calls.append(texts)
        if self.fail:
            raise ConnectionError("isolated provider failure")
        return [[1., 0.] if any(w in text.casefold() for w in ("tuition", "education costs", "fees")) else [0., 1.]
                for text in texts]


async def test_semantic_only_match_is_authorized_cited_and_falls_back_on_outage(hall):
    embedder = DeterministicEmbeddings()
    hall.indexer.embedder = embedder
    allowed, revision = hall.client.add("University payments", "Tuition is due on January 18.")
    denied, denied_revision = hall.client.add("Private payments", "Tuition is due on March 20.")
    hall.client.records[denied]["document"]["access_policy"] = {"allowed_agents": ["sentinel"]}
    await index_all(hall)
    assert any(j["kind"] == "library_embed" for j in hall.client.queued.values())
    await hall.indexer.embed_revision(allowed, revision)
    await hall.indexer.embed_revision(denied, denied_revision)
    result = await hall.service.search(hall.context, "education costs", limit=1)
    assert result.search_mode == "hybrid_rrf"
    assert result.results[0].document_id == allowed
    assert "January 18" in (await hall.service.citation_source(result.results[0].citation_id))["content"]
    embedder.fail = True
    result = await hall.service.search(hall.context, "Tuition")
    assert result.search_mode == "keyword_bm25" and result.results[0].document_id == allowed
    assert any("unavailable" in warning for warning in result.warnings)
    embedder.fail = False
    hall.client.records[allowed]["document"]["access_policy"] = {"allowed_agents": []}
    calls = len(embedder.calls)
    assert not (await hall.service.search(hall.context, "education costs")).results
    assert len(embedder.calls) == calls  # No query is sent when no authorized corpus remains.


async def test_semantic_missing_vectors_wrong_model_and_scope_limit_are_explicit(hall):
    embedder = DeterministicEmbeddings()
    hall.indexer.embedder = embedder
    doc, rev = hall.client.add("Payments", "Tuition fees are due.")
    await index_all(hall)
    assert any("pending" in w for w in (await hall.service.search(hall.context, "tuition")).warnings)
    await hall.indexer.embed_revision(doc, rev)
    embedder.fingerprint = "isolated:semantic-v2"
    assert not (await hall.service.search(hall.context, "education costs")).results
    await hall.indexer.embed_revision(doc, rev)
    other, other_rev = hall.client.add("Second payments", "Tuition fees in February.")
    await index_all(hall)
    await hall.indexer.embed_revision(other, other_rev)
    hall.indexer.semantic_max_passages = 1
    result = await hall.service.search(hall.context, "tuition")
    assert result.search_mode == "keyword_bm25" and any("Narrow" in w for w in result.warnings)
    scoped = await hall.service.search(hall.context, "education costs", doc_ids=[doc])
    assert scoped.results[0].document_id == doc


async def test_embedding_restart_resumes_committed_batches(hall):
    doc, rev = hall.client.add("Book", "\n".join("Tuition " + "a" * 1991 for _ in range(40)))
    await index_all(hall)
    class Interrupted(DeterministicEmbeddings):
        async def embed(self, texts):
            if self.calls:
                raise ConnectionError("simulated process restart")
            return await super().embed(texts)
    hall.indexer.embedder = Interrupted()
    with pytest.raises(ConnectionError):
        await hall.indexer.embed_revision(doc, rev)
    async with hall.sessions() as db:
        assert len((await db.scalars(select(LibraryEmbedding))).all()) == 32
    recovered = DeterministicEmbeddings()
    hall.indexer.embedder = recovered
    await hall.indexer.embed_revision(doc, rev)
    assert len(recovered.calls[0]) == 8
    assert (await hall.service.read_passage(hall.context, doc)).status == "success"


async def test_ocr_transcription_is_labeled_and_citations_keep_review_warning(hall):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    original = io.BytesIO()
    writer.write(original)
    class OCR:
        def pdf_page(self, blob, number):
            return "Final exam: January 18"
    hall.indexer.ocr_factory = OCR
    doc, rev = hall.client.add("Scanned calendar", original.getvalue(), mime="application/pdf")
    await index_all(hall)
    result = await hall.service.search(hall.context, "exam")
    assert result.results[0].needs_review and result.results[0].page == 1
    assert result.warnings and "OCR" in result.results[0].section
    source = await hall.service.citation_source(result.results[0].citation_id)
    assert source["needs_review"]
    assert (await hall.service.list_documents(hall.context)).documents[0].index_status == "ready_ocr"


def test_local_ocr_command_limits_cleanup_and_missing_runtime(tmp_path, monkeypatch):
    import struct
    from pathlib import Path
    from app.services.library_ocr import LocalOCR
    monkeypatch.setattr("app.services.library_ocr.shutil.which", lambda name: name)
    called, files = [], []
    def run(command, **kwargs):
        called.append(command)
        assert kwargs["timeout"] <= 300 and kwargs["stdin"] is not None
        image = Path(command[1])
        files.append(image)
        assert image.exists()
        Path(command[2] + ".txt").write_text("Observed OCR text", encoding="utf-8")
    monkeypatch.setattr("app.services.library_ocr.subprocess.run", run)
    png = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\x0dIHDR" + struct.pack(">II", 100, 100)
    assert LocalOCR().image(png, "image/png") == "Observed OCR text"
    assert called[0][3:] == ["-l", "tur+eng", "--psm", "3"]
    assert not files[0].exists()
    oversized = png[:16] + struct.pack(">II", 10000, 10000)
    with pytest.raises(ExtractionError, match="pixels"):
        LocalOCR().image(oversized, "image/png")
    monkeypatch.setattr("app.services.library_ocr.shutil.which", lambda name: None)
    with pytest.raises(ExtractionError, match="runtime"):
        LocalOCR().image(png, "image/png")
    with pytest.raises(ValueError):
        LocalOCR(languages="../secret")


def test_semantic_vectors_reject_invalid_or_unbounded_values():
    from app.services.library_embeddings import normalized
    for vector in ([], [0., 0.], [float("nan")], [float("inf")], [1.] * 4097):
        with pytest.raises(ValueError):
            normalized(vector)
    assert normalized([3., 4.]) == pytest.approx([0.6, 0.8])


async def test_changed_extraction_cannot_overwrite_published_citation(hall):
    doc, rev = hall.client.add("Original", "Exact evidence")
    await index_all(hall)
    citation = (await hall.service.read_passage(hall.context, doc)).citation_ids[0]
    async with hall.sessions() as db:
        generation = await db.get(LibraryIndexGeneration, hall.indexer.generation_id(rev))
        generation.status = "pending"
        await db.commit()
    from app.services.library_extract import ExtractedPassage
    import unittest.mock
    with unittest.mock.patch("app.services.library_indexer.extract_document_bytes", return_value=[ExtractedPassage("Changed transcription")]):
        with pytest.raises(ValueError, match="Published extraction changed"):
            await hall.indexer.index_revision(doc, rev)
    assert (await hall.service.citation_source(citation))["content"] == "Exact evidence"
    assert not (await hall.service.search(hall.context, "Changed")).results


def test_local_pdf_renderer_produces_bounded_page_and_cleans_scratch(monkeypatch):
    import shutil
    from app.services.library_ocr import LocalOCR, image_dimensions
    if not shutil.which("pdftoppm"):
        pytest.skip("Poppler is not installed on this test host")
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    original = io.BytesIO()
    writer.write(original)
    observed = []
    def recognize(self, image, directory):
        observed.append(image)
        width, height = image_dimensions(image.read_bytes(), "image/png")
        assert max(width, height) <= 2400
        return "isolated transcription fixture"
    monkeypatch.setattr(LocalOCR, "_recognize", recognize)
    ocr = LocalOCR()
    ocr.tesseract = "test-runtime"
    assert ocr.pdf_page(original.getvalue(), 1) == "isolated transcription fixture"
    assert observed and not observed[0].exists()


async def test_embedding_transport_restores_provider_input_order_and_rejects_missing_item(monkeypatch):
    from app.services.library_embeddings import LibraryEmbeddings
    class Transport:
        def __init__(self, **kwargs):
            self.embeddings = self
            self.data = [SimpleNamespace(index=1, embedding=[0., 2.]), SimpleNamespace(index=0, embedding=[3., 0.])]
        async def create(self, **kwargs):
            assert kwargs["input"] == ["first", "second"]
            return SimpleNamespace(data=self.data)
        async def close(self):
            self.closed = True
    monkeypatch.setattr("app.services.library_embeddings.AsyncOpenAI", Transport)
    embedder = LibraryEmbeddings(api_key="isolated", model="test-model")
    assert await embedder.embed(["first", "second"]) == [[1., 0.], [0., 1.]]
    embedder.client.data = [SimpleNamespace(index=1, embedding=[0., 2.])]
    with pytest.raises(ValueError, match="omitted"):
        await embedder.embed(["first", "second"])
    await embedder.aclose()
    assert embedder.client.closed
